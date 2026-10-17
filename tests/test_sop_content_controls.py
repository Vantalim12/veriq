import pytest

docx = pytest.importorskip("docx")

from docx.oxml import OxmlElement

from veriq.ingestion import InputError
from veriq.sop import read_sop


@pytest.mark.parametrize("location", ["inline", "body", "table"])
def test_content_controls_require_flattening_instead_of_losing_thresholds(tmp_path, location):
    document = docx.Document()
    paragraph = document.add_paragraph("Purchases above PHP ")
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = "10,000"
    run.append(value)
    if location == "inline":
        content.append(run)
        paragraph._p.append(control)
        paragraph.add_run(" must receive manager approval.")
        assert paragraph.text == "Purchases above PHP  must receive manager approval."
    else:
        nested_paragraph = OxmlElement("w:p")
        nested_paragraph.append(run)
        content.append(nested_paragraph)
        parent = document.element.body if location == "body" else document.add_table(rows=1, cols=1).cell(0, 0)._tc
        parent.append(control)
    control.append(content)
    path = tmp_path / "content-control.docx"
    document.save(path)

    with pytest.raises(InputError, match="content controls.*flattened"):
        read_sop(path)
