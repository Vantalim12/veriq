from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
import yaml

pymupdf = pytest.importorskip("pymupdf")
docx = pytest.importorskip("docx")

from veriq.cli import main
from veriq.ingestion import InputError, MAX_BYTES, read_bytes, read_policy, read_workflow
from veriq.reporting import markdown, write_artifact
from veriq.sop import (Draft, Review, approve_draft, document_reference, draft_requirements,
                       read_draft, read_review, read_sop, validate_span)
from veriq.verifier import verify

ROOT = Path(__file__).resolve().parents[1]
QUOTE = "Purchases above PHP 10,000 must receive manager approval before fulfillment."


@pytest.fixture(params=["pdf", "docx"])
def sop_path(request, tmp_path):
    path = tmp_path / ("sop." + request.param)
    if request.param == "pdf":
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_textbox(pymupdf.Rect(40, 40, 550, 200), "Procurement\n" + QUOTE)
            document.save(path)
    else:
        document = docx.Document()
        document.add_heading("Procurement", 1)
        document.add_paragraph(QUOTE)
        document.save(path)
    return path


def review_for(draft):
    return Review.model_validate({"schema_version": 1, "draft_hash": draft.content_hash, "policy_id": "procurement-sop", "policy_version": 1,
        "fields": {"body.amount": "number"}, "approved": {"reviewer": "human-reviewer", "reviewed_on": "2026-10-16"},
        "decisions": [{"candidate_id": candidate.id, "decision": "approve", "reason": "Reviewed the condition and control mapping",
                       "rule": {"id": f"REQ-{i}", "title": "Approval before fulfillment", "type": "must_pass_through",
                                "severity": "high", "scope": {"source_node_id": "intake", "target_node_id": "fulfillment"},
                                "gate_node_id": "approval-gate", "when": {"field": "body.amount", "operator": "gt", "value": 10000}}}
                      for i, candidate in enumerate(draft.candidates)]})


def test_located_extraction_and_repeatable_drafts(sop_path):
    document = read_sop(sop_path)
    assert document == read_sop(sop_path)
    draft = draft_requirements(document)
    assert draft == draft_requirements(document)
    assert draft.status == "draft" and len(draft.candidates) == 1
    source = draft.candidates[0].source
    assert source.quote.endswith(QUOTE)
    assert validate_span(source.model_dump(), document) == source
    assert (source.page == 1 and source.paragraph is None) if sop_path.suffix == ".pdf" else (source.paragraph == 2 and source.page is None and source.section == "Procurement")
    assert draft.candidates[0].actor is None  # Interpretation awaits human review.
    with pytest.raises(Exception):
        from veriq.contracts import Policy
        Policy.model_validate(draft.model_dump())


@pytest.mark.parametrize("field,value", [("quote", "invented"), ("page", 99), ("paragraph", 99), ("section", "invented"),
                                         ("block_id", "absent"), ("start", 99999), ("document_hash", "sha256:" + "0" * 64),
                                         ("extraction_version", "future")])
def test_citations_reject_invented_or_stale_spans(sop_path, field, value):
    document = read_sop(sop_path)
    span = draft_requirements(document).candidates[0].source.model_dump()
    span[field] = value
    with pytest.raises(InputError):
        validate_span(span, document)


def test_review_to_traceable_regression(sop_path, tmp_path):
    document = read_sop(sop_path)
    draft = draft_requirements(document)
    policy = approve_draft(draft, review_for(draft), document)
    rules = tmp_path / "rules.yaml"
    rules.write_text(yaml.safe_dump(policy.model_dump(exclude_none=True)))
    policy, policy_hash = read_policy(rules)
    baseline = read_workflow(ROOT / "examples/purchase-v1.json")
    candidate = read_workflow(ROOT / "examples/purchase-v2.json")
    audit = verify(candidate, policy, policy_hash, baseline, document=document, strict_unknown=True)
    assert audit["summary"]["exit_code"] == 1
    finding = audit["findings"][0]
    assert finding["source"]["validated"] is True
    assert finding["regression"] == "INTRODUCED"
    assert audit["traceability"][0]["evidence"]["counterexample_path"][-1] == "fulfillment"
    assert "SOP traceability" in markdown(audit)
    assert "validated" in markdown(audit)
    good = verify(baseline, policy, policy_hash, document=document, strict_unknown=True)
    missing = verify(baseline, policy, policy_hash, strict_unknown=True)
    assert good["summary"]["exit_code"] == 0
    assert missing["summary"]["exit_code"] == 3
    assert good["audit_id"] != missing["audit_id"]
    changed = deepcopy(document)
    changed["document_hash"] = "sha256:" + "0" * 64
    with pytest.raises(InputError):
        verify(candidate, policy, policy_hash, document=changed)


@pytest.mark.parametrize("mutation", ["missing", "duplicate", "unknown", "override_source", "override_reviewer", "reject_with_rule", "all_review_only"])
def test_review_boundary(sop_path, mutation):
    document = read_sop(sop_path)
    draft = draft_requirements(document)
    review = review_for(draft)
    decision = review.decisions[0]
    if mutation == "missing":
        review.decisions = []
    elif mutation == "duplicate":
        review.decisions.append(decision)
    elif mutation == "unknown":
        decision.candidate_id = "CAND-absent"
    elif mutation == "override_source":
        decision.rule["source"] = {}
    elif mutation == "override_reviewer":
        decision.rule["approved"] = {}
    elif mutation == "reject_with_rule":
        decision.decision = "reject"
    else:
        decision.decision = "review_only"
        decision.rule = None
    with pytest.raises(InputError):
        approve_draft(draft, review, document)


@pytest.mark.parametrize("change", ["source", "interpretation"])
def test_review_rejects_changed_draft_with_the_same_candidate_ids(sop_path, tmp_path, change):
    document = read_sop(sop_path)
    draft = draft_requirements(document)
    review = review_for(draft)
    if change == "source":
        updated = docx.Document()
        updated.add_paragraph(QUOTE.replace("10,000", "5,000"))
        path = tmp_path / "updated.docx"
        updated.save(path)
        document = read_sop(path)
        draft = draft_requirements(document)
    else:
        draft.candidates[0].interpretation = "Approval is optional."
    assert draft.candidates[0].id == review.decisions[0].candidate_id
    with pytest.raises(InputError, match="different draft"):
        approve_draft(draft, review, document)


def test_review_requires_draft_identity(sop_path):
    draft = draft_requirements(read_sop(sop_path))
    legacy = review_for(draft).model_dump()
    del legacy["draft_hash"]
    with pytest.raises(ValueError, match="draft_hash"):
        Review.model_validate(legacy)


def test_review_hash_survives_json_formatting(sop_path, tmp_path):
    document = read_sop(sop_path)
    draft = draft_requirements(document)
    review = review_for(draft)
    path = tmp_path / "reformatted.json"
    path.write_text(json.dumps(draft.model_dump(), sort_keys=True, ensure_ascii=True))
    reloaded = read_draft(path)
    assert reloaded.content_hash == draft.content_hash
    assert approve_draft(reloaded, review, document).rules[0].source["quote"] == draft.candidates[0].source.quote


def test_committed_example_review_matches_its_draft():
    document = read_sop(ROOT / "examples/procurement-sop.docx")
    draft = draft_requirements(document)
    review = read_review(ROOT / "examples/procurement-review.json")
    assert approve_draft(draft, review, document).rules[0].id == "REQ-004"


def test_cli_authoring_and_traceability(sop_path, tmp_path):
    draft_path, review_path, policy_path = [tmp_path / n for n in ("draft.json", "review.json", "rules.yaml")]
    assert main(["rules", "draft", "--sop", str(sop_path), "--output", str(draft_path)]) == 0
    draft = Draft.model_validate_json(draft_path.read_text())
    review_path.write_text(review_for(draft).model_dump_json())
    assert main(["rules", "approve", "--sop", str(sop_path), "--draft", str(draft_path), "--review", str(review_path), "--output", str(policy_path)]) == 0
    output = tmp_path / "reports"
    assert main(["verify", "--workflow", str(ROOT / "examples/purchase-v1.json"), "--rules", str(policy_path),
                 "--sop", str(sop_path), "--strict-unknown", "--output-dir", str(output)]) == 0
    assert json.loads((output / "audit.json").read_text())["traceability"][0]["source"]["validated"]
    before = draft_path.read_bytes()
    assert main(["rules", "draft", "--sop", str(sop_path), "--output", str(draft_path)]) == 2
    assert draft_path.read_bytes() == before
    assert draft_path.stat().st_mode & 0o077 == 0


def test_pdf_rejects_scanned_encrypted_and_excess_pages(tmp_path):
    for kind in ("empty", "encrypted", "long"):
        path = tmp_path / (kind + ".pdf")
        with pymupdf.open() as document:
            for _ in range(31 if kind == "long" else 1):
                page = document.new_page()
                if kind != "empty":
                    page.insert_text((50, 50), QUOTE)
            document.save(path, **({"encryption": pymupdf.PDF_ENCRYPT_AES_256, "owner_pw": "owner", "user_pw": "user"} if kind == "encrypted" else {}))
        with pytest.raises(InputError):
            read_sop(path)


def test_docx_tables_are_located_and_archive_attacks_rejected(tmp_path, monkeypatch):
    path = tmp_path / "table.docx"
    document = docx.Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = QUOTE
    document.save(path)
    assert draft_requirements(read_sop(path)).candidates[0].source.paragraph == 1
    original = path.read_bytes()
    for name, content in [("word/vbaProject.bin", b"macro"), ("word/evil.xml", b'<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a/>')]:
        with ZipFile(BytesIO(original)) as source, ZipFile(path, "w") as destination:
            for entry in source.infolist():
                destination.writestr(entry, source.read(entry))
            destination.writestr(name, content)
        with pytest.raises(InputError):
            read_sop(path)
    path.write_bytes(original)
    monkeypatch.setattr("veriq.sop.MAX_EXPANDED_BYTES", 100)
    with pytest.raises(InputError, match="decompression"):
        read_sop(path)


def test_model_quotes_are_validated_and_never_approved(sop_path, monkeypatch):
    document = read_sop(sop_path)
    block = document["blocks"][-1]
    monkeypatch.setattr("veriq.local_model.propose", lambda *args: [{"block_id": block["id"], "quote": QUOTE, "interpretation": "Edited draft meaning"}])
    draft = draft_requirements(document, "test-model")
    assert draft.candidates[0].status == "draft"
    assert draft.candidates[0].interpretation == "Edited draft meaning"
    monkeypatch.setattr("veriq.local_model.propose", lambda *args: [{"block_id": block["id"], "quote": "invented", "interpretation": "injection"}])
    with pytest.raises(InputError):
        draft_requirements(document, "test-model")


def test_artifact_protects_input_and_symlink(tmp_path):
    source = tmp_path / "input"
    source.write_text("keep")
    link = tmp_path / "output"
    link.symlink_to(source)
    with pytest.raises(ValueError):
        write_artifact("replace", link, [source])
    with pytest.raises(FileExistsError):
        write_artifact("replace", link, [])
    assert source.read_text() == "keep"


def test_authoring_artifact_uses_the_reader_byte_limit(tmp_path):
    payload = "界" * (MAX_BYTES // 3) + "x" * (MAX_BYTES % 3)
    accepted = tmp_path / "accepted"
    write_artifact(payload, accepted, [])
    assert read_bytes(accepted).decode("utf-8") == payload
    rejected = tmp_path / "new-directory/rejected"
    with pytest.raises(InputError, match="2 MiB input limit"):
        write_artifact(payload + "x", rejected, [])
    assert not rejected.parent.exists()


@pytest.mark.parametrize("command", ["draft", "approve"])
def test_cli_rejects_oversized_authoring_artifacts(tmp_path, capsys, command):
    document = docx.Document()
    document.add_heading("Section " + "x" * 24_000, 1)
    for index in range(100):
        document.add_paragraph(f"Control {index} must be checked.")
    source = tmp_path / "large-sections.docx"
    document.save(source)
    draft = draft_requirements(read_sop(source))
    assert len(draft.model_dump_json(indent=2).encode("utf-8")) > MAX_BYTES
    output = tmp_path / "output/artifact"
    arguments = ["rules", command, "--sop", str(source), "--output", str(output)]
    if command == "approve":
        # Rule titles and approval metadata can push the policy beyond its readable draft's size.
        draft_path, review_path = tmp_path / "draft.json", tmp_path / "review.json"
        draft.candidates = draft.candidates[:84]
        draft_path.write_text(draft.model_dump_json())
        review = review_for(draft)
        for decision in review.decisions:
            decision.rule["title"] = "T" * 512
        review_path.write_text(review.model_dump_json())
        assert draft_path.stat().st_size <= MAX_BYTES
        arguments.extend(["--draft", str(draft_path), "--review", str(review_path)])
    assert main(arguments) == 2
    assert "2 MiB input limit" in capsys.readouterr().err
    assert not output.parent.exists()


def test_document_limits_and_partial_review(sop_path, monkeypatch):
    document = read_sop(sop_path)
    draft = draft_requirements(document)
    extra = draft.candidates[0].model_copy(update={"id": "CAND-002"})
    draft.candidates.append(extra)
    review = review_for(draft)
    review.decisions[1].decision = "review_only"
    review.decisions[1].rule = None
    policy = approve_draft(draft, review, document)
    assert len(policy.rules) == 1
    document["document_hash"] = "sha256:" + "0" * 64
    with pytest.raises(InputError, match="different document"):
        approve_draft(draft, review, document)
    monkeypatch.setattr("veriq.sop.MAX_TEXT", 10)
    with pytest.raises(InputError, match="text limit"):
        read_sop(sop_path)


def test_ollama_loopback_transport_validation(monkeypatch):
    from veriq.local_model import propose
    sent = {}
    response = {"done": True, "response": json.dumps({"candidates": []})}

    class Connection:
        def __init__(self, host, port, timeout):
            assert (host, port, timeout) == ("127.0.0.1", 11434, 45)
        def request(self, method, path, payload, headers):
            sent.update(json.loads(payload))
            assert method == "POST" and path == "/api/generate"
        def getresponse(self):
            return self
        status = 200
        def read(self, size):
            return json.dumps(response).encode()
        def close(self):
            sent["closed"] = True

    monkeypatch.setattr("veriq.local_model.http.client.HTTPConnection", Connection)
    assert propose({"blocks": []}, "test-model") == []
    assert sent["stream"] is False and sent["format"] == "json" and sent["closed"]
    for invalid in ({"done": False}, {"done": True, "response": "invalid"},
                    {"done": True, "response": json.dumps({"approved": True})},
                    {"done": True, "done_reason": "length", "response": "{}"}):
        response.clear()
        response.update(invalid)
        with pytest.raises(InputError):
            propose({"blocks": []}, "test-model")


def test_heuristics_preserve_decimal_thresholds_and_wrapped_conditions():
    value = "Purchases above 10.50 must receive approval\nbefore fulfillment unless explicitly exempt."
    document = {"document_hash": "sha256:" + "0" * 64, "format": "pdf", "extraction_version": "sop-1",
                "parser": "test", "limitations": [], "blocks": [{"id": "page-1", "page": 1, "paragraph": None,
                "section": None, "text": value}]}
    draft = draft_requirements(document)
    assert len(draft.candidates) == 1
    assert draft.candidates[0].source.quote == value


@pytest.mark.parametrize("field_tag", ["fldSimple", "fldChar", "instrText"])
def test_docx_fields_require_flattening_instead_of_losing_thresholds(tmp_path, field_tag):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()
    paragraph = document.add_paragraph("Purchases above PHP ")
    field = OxmlElement("w:" + field_tag)
    if field_tag == "fldSimple":
        field.set(qn("w:instr"), "REF ApprovalLimit")
        run = OxmlElement("w:r")
        value = OxmlElement("w:t")
        value.text = "10,000"
        run.append(value)
        field.append(run)
        paragraph._p.append(field)
    else:
        paragraph.add_run()._r.append(field)
        paragraph.add_run("10,000")
    paragraph.add_run(" must receive approval before fulfillment.")
    path = tmp_path / "field.docx"
    document.save(path)
    with pytest.raises(InputError, match="fields.*flattened"):
        read_sop(path)


def test_repeated_requirements_keep_distinct_source_positions(tmp_path, monkeypatch):
    sentence = "Managers must approve."
    document = docx.Document()
    document.add_paragraph("  " + sentence + " \n " + sentence + "  ")
    path = tmp_path / "repeated.docx"
    document.save(path)
    extracted = read_sop(path)
    draft = draft_requirements(extracted)
    assert len(draft.candidates) == 2
    first, second = [candidate.source for candidate in draft.candidates]
    assert first.quote == second.quote == sentence
    assert first.start == 2 and second.start > first.end
    for candidate in draft.candidates:
        validate_span(candidate.source.model_dump(), extracted)
    monkeypatch.setattr("veriq.local_model.propose", lambda *args: [
        {"block_id": first.block_id, "quote": sentence, "interpretation": sentence}])
    with pytest.raises(InputError, match="ambiguous"):
        draft_requirements(extracted, "test-model")
